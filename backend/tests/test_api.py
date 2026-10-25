import io
import time
import pytest
from fastapi.testclient import TestClient
from pypdf import PdfWriter
import docx

from app.main import app
from app.services.session_manager import session_manager, Session
from app.services.document_parser import extract_document_text, DocumentParsingError
from app.models.schemas import DocumentAnalysis, RiskClause

client = TestClient(app)


def test_health_endpoint():
    response = client.get("/api/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    assert "providers" in data
    assert "session" in data
    assert "openai" in data["providers"]


def test_languages_endpoints():
    response = client.get("/api/languages")
    assert response.status_code == 200
    data = response.json()
    assert data["count"] == 16
    lang_codes = [lang["code"] for lang in data["languages"]]
    assert "en" in lang_codes
    assert "hi" in lang_codes
    assert "ta" in lang_codes
    assert "ur" in lang_codes
    assert "sa" in lang_codes
    assert "kok" in lang_codes

    # Capabilities endpoint
    cap_resp = client.get("/api/languages/capabilities")
    assert cap_resp.status_code == 200
    caps = cap_resp.json()["matrix"]
    # Check that Urdu is marked RTL
    assert caps["ur"]["direction"] == "rtl"
    # Check that Konkani speech is truthfully reported as false for cloud STT
    assert caps["kok"]["capabilities"]["stt_whisper"] is False


def test_document_parser_txt():
    content = "This is an Indian Residential Rental Agreement between Landlord and Tenant."
    text, meta = extract_document_text("rental.txt", content.encode("utf-8"))
    assert text == content
    assert "encoding" in meta


def test_document_parser_docx():
    doc = docx.Document()
    doc.add_paragraph("Clause 1: The tenant shall pay a monthly rent of Rs. 25,000.")
    doc.add_paragraph("Clause 2: Notice period is 30 days.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Security Deposit"
    table.rows[0].cells[1].text = "Rs. 1,00,000"

    bio = io.BytesIO()
    doc.save(bio)
    file_bytes = bio.getvalue()

    text, meta = extract_document_text("lease.docx", file_bytes)
    assert "Clause 1" in text
    assert "Security Deposit" in text
    assert meta["paragraph_count"] >= 2


def test_document_parser_pdf():
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    # Blank page without text should trigger scanned/empty check
    bio = io.BytesIO()
    writer.write(bio)

    with pytest.raises(DocumentParsingError) as exc_info:
        extract_document_text("blank.pdf", bio.getvalue())
    assert exc_info.value.code in ["SCANNED_DOCUMENT_OR_IMAGE_PDF", "EMPTY_DOCUMENT"]


def test_document_parser_empty_file():
    with pytest.raises(DocumentParsingError) as exc_info:
        extract_document_text("empty.txt", b"")
    assert exc_info.value.code == "EMPTY_FILE"


def test_document_parser_unsupported_type():
    with pytest.raises(DocumentParsingError) as exc_info:
        extract_document_text("malicious.exe", b"binarycontent")
    assert exc_info.value.code == "UNSUPPORTED_FILE_TYPE"


def test_analyze_missing_api_key_honest_error(monkeypatch):
    # If API key is not configured, must return 503 PROVIDER_NOT_CONFIGURED, never fake data!
    monkeypatch.setattr("app.config.settings.OPENAI_API_KEY", "your_openai_api_key_here")
    monkeypatch.setattr("app.config.settings.GEMINI_API_KEY", "your_gemini_api_key_here")
    content = "Rental Agreement: The tenant agrees to pay 50,000 deposit and 15,000 monthly rent."
    files = {"file": ("contract.txt", content.encode("utf-8"), "text/plain")}
    response = client.post("/api/analyze", files=files, data={"language": "en"})

    assert response.status_code == 503
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] == "PROVIDER_NOT_CONFIGURED"
    assert "temporarily unavailable" in data["error"]["message"].lower()


def test_session_manager_lifecycle_and_deletion():
    sm = session_manager
    session = sm.get_or_create_session()
    sid = session.session_id

    # Add document
    doc = sm.add_document(
        session_id=sid,
        document_id="doc_123",
        filename="test.txt",
        file_size=100,
        content_type="text/plain",
        extracted_text="Test legal content"
    )
    assert doc["document_id"] == "doc_123"
    assert sm.get_document(sid, "doc_123") is not None

    # Delete single document
    del_resp = client.delete(f"/api/documents/doc_123?session_id={sid}")
    assert del_resp.status_code == 200
    assert sm.get_document(sid, "doc_123") is None

    # Clear entire session
    clear_resp = client.post(f"/api/session/clear?session_id={sid}")
    assert clear_resp.status_code == 200
    assert sm.get_session(sid) is None


def test_session_expiration_cleanup():
    sm = session_manager
    session = sm.get_or_create_session("expired_session_test")
    # Simulate past timestamp beyond TTL
    session.last_accessed_at = time.time() - (35 * 60)

    cleaned = sm.cleanup_expired_sessions()
    assert cleaned >= 1
    assert sm.get_session("expired_session_test") is None


def test_chat_without_document(monkeypatch):
    # Chatting without a document should provide general legal information or transparent missing key
    monkeypatch.setattr("app.config.settings.OPENAI_API_KEY", "your_openai_api_key_here")
    monkeypatch.setattr("app.config.settings.GEMINI_API_KEY", "your_gemini_api_key_here")
    payload = {
        "message": "What is an indemnity clause in an Indian commercial contract?",
        "language": "en"
    }
    response = client.post("/api/chat", json=payload)
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "PROVIDER_NOT_CONFIGURED"


def test_unsupported_speech_language_honest_rejection():
    # Uploading speech in Konkani (kok) which is truthfully not supported by cloud speech engines
    audio_content = b"RIFF....WAVEfmt ...."
    files = {"file": ("sample.wav", audio_content, "audio/wav")}
    response = client.post("/api/speech/transcribe", files=files, data={"language": "kok"})
    assert response.status_code in [400, 503]
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] in ["UNSUPPORTED_SPEECH_LANGUAGE", "PROVIDER_NOT_CONFIGURED"]


def test_translation_missing_provider_honest_error(monkeypatch):
    monkeypatch.setattr("app.config.settings.OPENAI_API_KEY", "your_openai_api_key_here")
    monkeypatch.setattr("app.config.settings.GEMINI_API_KEY", "your_gemini_api_key_here")
    monkeypatch.setattr("app.config.settings.GOOGLE_TRANSLATION_API_KEY", "your_google_translation_api_key_here")
    payload = {
        "text": "The Lessee shall not sublet or assign the leased premises.",
        "target_language": "hi",
        "scope": "selected_clause"
    }
    response = client.post("/api/translate", json=payload)
    assert response.status_code == 503
    data = response.json()
    assert data["error"]["code"] == "PROVIDER_NOT_CONFIGURED"


def test_analyze_document_mock_success(monkeypatch):
    mock_data = DocumentAnalysis(
        document_title="Sample Lease Agreement",
        document_type="Residential Lease",
        detected_language="English",
        summary="An 11-month lease agreement in Bengaluru.",
        key_obligations=["Pay rent of 25,000 INR on the 1st of each month."],
        important_dates=["Duration: 11 months", "Lock-in period: 3 months"],
        financial_terms=["Security deposit: 100,000 INR"],
        risks=[
            RiskClause(
                id="risk-1",
                clause_title="Maintenance Fee",
                original_text="Non-refundable maintenance charge of 5,000 INR",
                risk_level="low",
                reason="Standard society maintenance charge.",
                potential_consequences="Tenant pays additional fee monthly.",
                suggested_precautions="Verify society receipts."
            )
        ],
        precautions=["Verify ownership before signing."],
        questions_for_lawyer=["Is 11-month lease registration mandatory?"],
        limitations="Analysis based only on the provided text."
    )
    monkeypatch.setattr("app.api.routes.analyze.analyze_document_text", lambda **kwargs: mock_data)
    files = {"file": ("lease.txt", b"Lease Agreement text content", "text/plain")}
    response = client.post("/api/analyze", files=files, data={"language": "en"})
    assert response.status_code == 200
    data = response.json()
    assert data["analysis"]["document_title"] == "Sample Lease Agreement"
    assert "disclaimer" in data
    assert "document_id" in data


def test_chat_mock_success(monkeypatch):
    monkeypatch.setattr(
        "app.api.routes.chat.chat_with_legal_assistant",
        lambda **kwargs: {
            "response": "An indemnity clause protects one party against legal liabilities or losses.",
            "is_general_legal_info": True
        }
    )
    payload = {
        "message": "What is an indemnity clause?",
        "language": "en"
    }
    response = client.post("/api/chat", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert "indemnity" in data["response"]
    assert data["is_general_legal_info"] is True
    assert "disclaimer" in data


def test_translation_mock_success(monkeypatch):
    async def mock_translate(**kwargs):
        return {
            "translated_text": "पट्टाकर्ता परिसर को खाली नहीं करेगा।",
            "source_language": "en",
            "target_language": "hi",
            "provider": "mock_openai",
            "scope": "selected_clause"
        }
    monkeypatch.setattr("app.api.routes.translate.translate_text", mock_translate)
    payload = {
        "text": "The Lessee shall not sublet or assign the leased premises.",
        "target_language": "hi",
        "scope": "selected_clause"
    }
    response = client.post("/api/translate", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert data["translated_text"] == "पट्टाकर्ता परिसर को खाली नहीं करेगा।"
    assert data["target_language"] == "hi"


def test_free_google_translation_live():
    # Direct test of free translation API without requiring server keys
    payload = {
        "text": "This Agreement shall be governed by the laws of India.",
        "target_language": "hi",
        "scope": "selected_clause",
        "provider": "free_google"
    }
    response = client.post("/api/translate", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert "translated_text" in data
    assert len(data["translated_text"]) > 0
    assert "Fast Legal Translation" in data["provider"]


