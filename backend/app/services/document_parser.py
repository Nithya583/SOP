import io
import os
from typing import Tuple, Dict, Any, Optional
from pypdf import PdfReader
from pypdf.errors import PdfReadError, FileNotDecryptedError
import docx
from app.config import settings


class DocumentParsingError(Exception):
    def __init__(self, code: str, message: str, suggestion: Optional[str] = None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.suggestion = suggestion


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
ALLOWED_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "text/plain",
    "application/octet-stream"  # Some browsers send octet-stream for docx/txt
}


def validate_file_metadata(filename: str, file_size: int, content_type: Optional[str] = None):
    if not filename:
        raise DocumentParsingError(
            code="INVALID_FILENAME",
            message="No filename provided.",
            suggestion="Please select a valid document file."
        )

    _, ext = os.path.splitext(filename.lower())
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentParsingError(
            code="UNSUPPORTED_FILE_TYPE",
            message=f"File extension '{ext}' is not supported. Only PDF, DOCX, and TXT files are accepted.",
            suggestion="Please convert your document to PDF, Word (DOCX), or Plain Text (TXT)."
        )

    if file_size <= 0:
        raise DocumentParsingError(
            code="EMPTY_FILE",
            message="Uploaded file is empty (0 bytes).",
            suggestion="Please check that you have selected the correct document file."
        )

    if file_size > settings.max_upload_size_bytes:
        raise DocumentParsingError(
            code="FILE_TOO_LARGE",
            message=f"File size exceeds the {settings.MAX_UPLOAD_SIZE_MB}MB limit.",
            suggestion=f"Please upload a document smaller than {settings.MAX_UPLOAD_SIZE_MB}MB."
        )


def extract_text_from_pdf(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
    except FileNotDecryptedError:
        raise DocumentParsingError(
            code="PASSWORD_PROTECTED_DOCUMENT",
            message="This PDF is password-protected and cannot be analyzed.",
            suggestion="Please remove the password from the PDF before uploading."
        )
    except Exception as e:
        raise DocumentParsingError(
            code="CORRUPTED_DOCUMENT",
            message=f"Could not read the PDF document: {str(e)}",
            suggestion="Ensure the PDF file is not corrupted or try resaving it."
        )

    if reader.is_encrypted:
        try:
            # Try decrypting with empty password
            decrypted = reader.decrypt("")
            if not decrypted:
                raise DocumentParsingError(
                    code="PASSWORD_PROTECTED_DOCUMENT",
                    message="This PDF is password-protected and cannot be analyzed.",
                    suggestion="Please remove the password from the PDF before uploading."
                )
        except Exception:
            raise DocumentParsingError(
                code="PASSWORD_PROTECTED_DOCUMENT",
                message="This PDF is password-protected and cannot be analyzed.",
                suggestion="Please remove the password from the PDF before uploading."
            )

    page_count = len(reader.pages)
    if page_count == 0:
        raise DocumentParsingError(
            code="EMPTY_DOCUMENT",
            message="The PDF contains no pages.",
            suggestion="Please provide a valid legal document."
        )

    extracted_pages = []
    for idx, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
            if text.strip():
                extracted_pages.append(text.strip())
        except Exception:
            continue

    full_text = "\n\n".join(extracted_pages).strip()

    # Detect scanned PDF / image-only PDF
    if len(full_text) < 50:
        raise DocumentParsingError(
            code="SCANNED_DOCUMENT_OR_IMAGE_PDF",
            message="The PDF appears to be a scanned image or photograph without selectable digital text.",
            suggestion="AI Legal Guardian requires text-based digital documents. Please use an OCR tool to convert scanned documents to searchable PDF or text, or upload a digital version."
        )

    return full_text, {"page_count": page_count, "extracted_pages": len(extracted_pages)}


def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise DocumentParsingError(
            code="CORRUPTED_DOCUMENT",
            message=f"Could not read the Word (DOCX) document: {str(e)}",
            suggestion="Ensure the document is a valid Microsoft Word .docx file."
        )

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Extract text from tables as well
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                table_texts.append(" | ".join(row_text))

    full_text = "\n\n".join(paragraphs + table_texts).strip()

    if not full_text:
        raise DocumentParsingError(
            code="EMPTY_DOCUMENT",
            message="The Word document does not contain any readable text.",
            suggestion="Please check that the document has content and is not empty."
        )

    return full_text, {"paragraph_count": len(paragraphs), "table_count": len(doc.tables)}


def extract_text_from_txt(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            text = file_bytes.decode(enc).strip()
            if not text:
                raise DocumentParsingError(
                    code="EMPTY_DOCUMENT",
                    message="The text file is empty.",
                    suggestion="Please provide a text file with content."
                )
            return text, {"encoding": enc}
        except UnicodeDecodeError:
            continue

    raise DocumentParsingError(
        code="UNREADABLE_TEXT_ENCODING",
        message="Could not decode the text file using standard encodings (UTF-8, Latin-1).",
        suggestion="Please save your file with UTF-8 encoding."
    )


def extract_document_text(filename: str, file_bytes: bytes, content_type: Optional[str] = None) -> Tuple[str, Dict[str, Any]]:
    validate_file_metadata(filename, len(file_bytes), content_type)
    _, ext = os.path.splitext(filename.lower())

    if ext == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext == ".docx":
        return extract_text_from_docx(file_bytes)
    elif ext == ".txt":
        return extract_text_from_txt(file_bytes)
    else:
        raise DocumentParsingError(
            code="UNSUPPORTED_FILE_TYPE",
            message=f"Unsupported file type: {ext}",
            suggestion="Please upload a PDF, DOCX, or TXT file."
        )
